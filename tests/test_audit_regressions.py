"""One regression test per defect found in the September 2026 audit.

Each test names the finding it pins. They exist so that a defect which once shipped
silently cannot return silently: every one of these failed before its fix and passes
after it.

The audit's theme was that all three critical defects produced a plausible,
confidently-labelled, wrong number with no error and no warning. Most of these tests
therefore assert on a *warning* or on agreement between two independently computed
figures, not merely that the code ran.
"""

from __future__ import annotations

import warnings
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.dcf.bridge import base_share_count, build_bridge
from src.dcf.engine import DCFEngine, closed_form_dilution_price, terminal_discount_factor
from src.dcf.projector import Projector
from src.dcf.sensitivity import _beta_for_wacc, football_field, wacc_vs_growth
from src.dcf.wacc import WACCCalculator
from src.excel.builder import build_excel_model
from src.fetcher.normalizer import FinancialNormalizer, to_float
from src.fetcher.yfinance_client import YFinanceClient
from src.models.assumptions import DCFAssumptions
from src.models.errors import DataQualityError, ValuationError
from src.models.financials import Financials, total_cash_position
from src.models.quality_gate import DataQualityGate
from tests.conftest import make_financials

P1, P2 = pd.Timestamp("2024-12-31"), pd.Timestamp("2025-12-31")


def _canon(income: dict, balance: dict, cashflow: dict, periods=(P1, P2)) -> pd.DataFrame:
    """Build canonical statements from raw provider-style row labels."""

    def frame(rows: dict) -> pd.DataFrame:
        return pd.DataFrame({p: rows for p in periods}).T.T

    return FinancialNormalizer.canonicalize(
        {"income": frame(income), "balance": frame(balance), "cashflow": frame(cashflow)}
    )


class TestC1ScenarioResolution:
    """Scenario overrides were silently dropped outside the repo root."""

    def test_scenarios_differ_from_any_working_directory(self, tmp_path, monkeypatch):
        monkeypatch.chdir(tmp_path)
        values = {}
        for scenario in ("bear", "base", "bull"):
            assumptions = DCFAssumptions.from_yaml(scenario=scenario)
            values[scenario] = assumptions.terminal.perpetuity_growth

        assert values["bear"] < values["base"] < values["bull"], (
            "scenario overrides were not applied; config paths must resolve against the "
            "package root, not the working directory"
        )

    def test_missing_scenarios_file_refuses_rather_than_falling_back(self, tmp_path):
        """Never return base-case numbers under a bull or bear heading."""
        assumptions_only = tmp_path / "assumptions.yaml"
        assumptions_only.write_text("projection:\n  years: 5\n", encoding="utf-8")

        with pytest.raises(FileNotFoundError, match="bull"):
            DCFAssumptions.from_yaml(
                assumptions_only, scenario="bull", scenarios_path=tmp_path / "nope.yaml"
            )

    def test_base_scenario_still_works_without_a_scenarios_file(self, tmp_path):
        assumptions_only = tmp_path / "assumptions.yaml"
        assumptions_only.write_text("projection:\n  years: 5\n", encoding="utf-8")
        loaded = DCFAssumptions.from_yaml(
            assumptions_only, scenario="base", scenarios_path=tmp_path / "nope.yaml"
        )
        assert loaded.projection.years == 5


class TestC2MixedVintage:
    """FY(N) income silently bridged against an FY(N-1) balance sheet."""

    def test_income_only_latest_period_warns(self):
        statements = FinancialNormalizer.canonicalize(
            {
                "income": pd.DataFrame(
                    {
                        P1: {"Total Revenue": 900.0, "EBIT": 180.0},
                        P2: {"Total Revenue": 1000.0, "EBIT": 200.0},
                    }
                ).T.T,
                "balance": pd.DataFrame(
                    {P1: {"Total Debt": 100.0, "Cash And Cash Equivalents": 50.0}}
                ).T.T,
                "cashflow": pd.DataFrame({P1: {"Operating Cash Flow": 220.0}}).T.T,
            }
        )
        financials = Financials(ticker="STUB", statements=statements, info={})
        report = DataQualityGate(financials).validate()
        assert any("no balance sheet" in w for w in report.warnings), (
            "a period carrying an income statement but no balance sheet must be flagged"
        )

    def test_individually_stale_field_warns(self):
        """AAPL's shipped fixture has no interest expense for its two latest years."""
        financials = YFinanceClient("AAPL", offline_mode=True).get_financials()
        report = DataQualityGate(financials).validate()
        assert any("earlier year is used" in w for w in report.warnings)


class TestC3ExcelCapitalStructure:
    """The workbook hardcoded market weights and ignored capital_structure='target'."""

    def test_weight_formulas_honour_the_target_switch(self, tmp_path):
        warnings.simplefilter("ignore")
        financials = YFinanceClient("AAPL", offline_mode=True).get_financials()
        assumptions = DCFAssumptions.from_yaml(
            overrides={"wacc": {"capital_structure": "target", "target_debt_weight": 0.20}}
        )
        result = DCFEngine(financials, assumptions, ticker="AAPL").run()
        path = build_excel_model(result, tmp_path / "t.xlsx", financials=financials)

        book = load_workbook(path)
        inputs_labels = [
            book["Inputs"].cell(row=r, column=1).value for r in range(1, book["Inputs"].max_row + 1)
        ]
        assert any(
            isinstance(v, str) and v.startswith("Capital structure") for v in inputs_labels
        ), "the workbook must expose the capital-structure switch as an input"

        wacc_sheet = book["WACC"]
        debt_weight_formula = None
        for r in range(1, wacc_sheet.max_row + 1):
            if wacc_sheet.cell(row=r, column=1).value == "Debt weight":
                debt_weight_formula = wacc_sheet.cell(row=r, column=2).value
        assert debt_weight_formula and "target" in debt_weight_formula, (
            "the debt-weight formula must branch on the capital-structure cell"
        )


class TestH1ShareCountBasis:
    """The denominator was last year's weighted average, not today's count."""

    def test_current_shares_outstanding_wins(self):
        financials = make_financials(
            revenue=1000.0,
            ebit=200.0,
            diluted_shares=100.0,
            ordinary_shares=110.0,
            info={"sharesOutstanding": 120.0},
        )
        assert base_share_count(financials) == pytest.approx(120.0)

    def test_balance_sheet_count_beats_the_weighted_average(self):
        financials = make_financials(diluted_shares=100.0, ordinary_shares=110.0)
        assert base_share_count(financials) == pytest.approx(110.0)

    def test_weighted_average_fallback_warns(self):
        financials = make_financials(diluted_shares=100.0)
        with pytest.warns(UserWarning, match="backward-looking"):
            assert base_share_count(financials) == pytest.approx(100.0)

    def test_tesla_uses_current_not_average(self):
        """The concrete case: 3.53bn average against 3.95bn actual, 12% of value."""
        financials = YFinanceClient("TSLA", offline_mode=True).get_financials()
        assert base_share_count(financials) == pytest.approx(
            float(financials.info["sharesOutstanding"])
        )
        assert base_share_count(financials) > financials.latest("diluted_shares")


class TestH2CashDoubleCount:
    """The combined cash row already includes short-term investments."""

    def test_combined_row_is_not_added_to_investments_again(self):
        statements = _canon(
            {"Total Revenue": 1000.0, "EBIT": 200.0},
            {
                "Cash Cash Equivalents And Short Term Investments": 500.0,
                "Other Short Term Investments": 200.0,
                "Total Debt": 100.0,
            },
            {"Operating Cash Flow": 250.0},
        )
        financials = Financials(ticker="X", statements=statements, info={"marketCap": 5000.0})
        assert total_cash_position(financials) == pytest.approx(500.0)
        assert WACCCalculator(financials).cash == pytest.approx(500.0)

    def test_separate_rows_still_sum(self):
        statements = _canon(
            {"Total Revenue": 1000.0, "EBIT": 200.0},
            {
                "Cash And Cash Equivalents": 300.0,
                "Other Short Term Investments": 200.0,
                "Total Debt": 100.0,
            },
            {"Operating Cash Flow": 250.0},
        )
        financials = Financials(ticker="X", statements=statements, info={})
        assert total_cash_position(financials) == pytest.approx(500.0)


class TestH3DebtFreeCompany:
    def test_no_debt_row_is_valued_with_a_warning(self):
        statements = _canon(
            {"Total Revenue": 1000.0, "EBIT": 200.0},
            {"Cash And Cash Equivalents": 60.0, "Ordinary Shares Number": 100.0},
            {"Operating Cash Flow": 250.0},
        )
        financials = Financials(
            ticker="NODEBT", statements=statements, info={"marketCap": 5000.0}
        )
        report = DataQualityGate(financials).validate()
        assert report.ok
        assert any("total_debt" in w for w in report.warnings)


class TestH5InterestExpenseSign:
    """abs() turned net interest income into an 8x cost of debt."""

    def test_negative_interest_is_not_flipped(self):
        statements = _canon(
            {"Total Revenue": 1000.0, "EBIT": 200.0, "Interest Expense": -40.0},
            {"Total Debt": 100.0, "Cash And Cash Equivalents": 900.0},
            {"Operating Cash Flow": 250.0},
        )
        assert statements.loc["interest_expense"].iloc[-1] == pytest.approx(-40.0)

        financials = Financials(ticker="X", statements=statements, info={"marketCap": 5000.0})
        calc = WACCCalculator(financials)
        assert calc.cost_of_debt == pytest.approx(calc.risk_free_rate), (
            "net interest income must fall back to the risk-free rate, not be charged"
        )

    def test_positive_interest_still_drives_cost_of_debt(self):
        statements = _canon(
            {"Total Revenue": 1000.0, "EBIT": 200.0, "Interest Expense": 5.0},
            {"Total Debt": 100.0, "Cash And Cash Equivalents": 50.0},
            {"Operating Cash Flow": 250.0},
        )
        financials = Financials(ticker="X", statements=statements, info={"marketCap": 5000.0})
        assert WACCCalculator(financials).cost_of_debt == pytest.approx(0.05)


class TestH6MissingShareCount:
    def test_raises_a_valuation_error_the_cli_can_catch(self):
        statements = _canon(
            {"Total Revenue": 1000.0, "EBIT": 200.0},
            {"Total Debt": 100.0, "Cash And Cash Equivalents": 60.0},
            {"Operating Cash Flow": 250.0},
        )
        financials = Financials(ticker="X", statements=statements, info={"marketCap": 5000.0})
        with pytest.raises(ValuationError):
            DCFEngine(financials, DCFAssumptions.from_yaml(), ticker="X").run()
        assert issubclass(DataQualityError, ValuationError)


class TestH7TerminalDiscounting:
    """An exit multiple is a sale price, not a flow stream."""

    def test_gordon_keeps_the_mid_year_lift(self):
        assert terminal_discount_factor(0.10, 5, True, "gordon") == pytest.approx(1 / 1.1**4.5)

    def test_exit_multiple_discounts_at_the_full_year(self):
        assert terminal_discount_factor(0.10, 5, True, "exit_multiple") == pytest.approx(
            1 / 1.1**5
        )

    def test_the_two_differ_by_exactly_the_half_year(self):
        gordon = terminal_discount_factor(0.10, 5, True, "gordon")
        exit_mult = terminal_discount_factor(0.10, 5, True, "exit_multiple")
        assert gordon / exit_mult == pytest.approx(1.1**0.5)


class TestH9NanGuard:
    """The old NaN test could not fail: openpyxl round-trips NaN to None.

    The replacement written during this audit could not fail either -- see
    `tests/test_second_audit_regressions.py`, which corrects it. The check now runs
    against the in-memory workbook, the only place a non-finite value is still visible.
    """

    def test_num_blanks_nan_and_refuses_infinity(self):
        """NaN means "not reported" and blanks. An infinity is a bug and raises.

        `_num` used to blank both. Blanking an infinity leaves a cell that every
        formula pointing at it reads as zero -- a wrong number presented with no
        warning, which is the failure mode this project exists to refuse.
        """
        from src.excel.builder import _num

        assert _num(float("nan")) is None
        assert _num(1.5) == 1.5
        for bad in (float("inf"), float("-inf")):
            with pytest.raises(ValueError, match="infinite"):
                _num(bad)

    def test_no_workbook_cell_holds_a_non_finite_float(self):
        from src.excel.builder import ExcelModelBuilder, assert_all_cells_finite

        warnings.simplefilter("ignore")
        financials = YFinanceClient("AAPL", offline_mode=True).get_financials()
        result = DCFEngine(financials, DCFAssumptions.from_yaml(), ticker="AAPL").run()

        book = ExcelModelBuilder(result=result, financials=financials).workbook()
        assert_all_cells_finite(book)  # raises, rather than skipping, on a leak


class TestH10BetaSolver:
    """_beta_for_wacc had no direct coverage; the monotonicity test did not bind."""

    @pytest.mark.parametrize("target", [0.06, 0.08, 0.10, 0.14])
    def test_solved_beta_reproduces_the_target_wacc(self, target):
        warnings.simplefilter("ignore")
        financials = YFinanceClient("AAPL", offline_mode=True).get_financials()
        assumptions = DCFAssumptions.from_yaml()
        base = DCFEngine(financials, assumptions, ticker="AAPL").run()

        beta = _beta_for_wacc(base, target)
        patched = assumptions.model_copy(deep=True)
        patched.wacc.beta_override = beta
        achieved = WACCCalculator(financials, patched).wacc
        assert achieved == pytest.approx(target, abs=1e-9)

    def test_sensitivity_rows_actually_span_different_waccs(self):
        """The grid must move with WACC, not merely be monotonic by accident."""
        warnings.simplefilter("ignore")
        financials = YFinanceClient("AAPL", offline_mode=True).get_financials()
        assumptions = DCFAssumptions.from_yaml()
        base = DCFEngine(financials, assumptions, ticker="AAPL").run()
        frame = wacc_vs_growth(financials, assumptions, base).frame

        first_col = frame.iloc[:, 0].dropna()
        assert first_col.iloc[0] / first_col.iloc[-1] > 1.2, (
            "a live WACC axis must produce materially different values top to bottom"
        )


class TestM1M2Buybacks:
    def test_closed_form_tracks_the_solver_under_buybacks(self):
        warnings.simplefilter("ignore")
        financials = YFinanceClient("AAPL", offline_mode=True).get_financials()
        for offset in (0.0, 0.5, 1.0):
            assumptions = DCFAssumptions.from_yaml(
                overrides={"sbc": {"method": "dilute", "buyback_offset_pct": offset}}
            )
            engine = DCFEngine(financials, assumptions, ticker="AAPL")
            result = engine.run()
            core = engine._compute_core()
            shares = base_share_count(financials)
            equity = build_bridge(core.enterprise_value, financials, assumptions, shares).equity_value
            sbc = [float(v) for v in core.projection.table.loc["sbc"]]

            expected = closed_form_dilution_price(
                equity, shares, sbc, core.wacc_calc.cost_of_equity, offset
            )
            assert result.value_per_share == pytest.approx(expected, rel=2e-3), (
                f"closed form and solver disagree at buyback_offset_pct={offset}"
            )

    def test_full_buyback_equals_expensing(self):
        """Issue stock then buy it all back: economically identical to expensing it."""
        warnings.simplefilter("ignore")
        financials = YFinanceClient("AAPL", offline_mode=True).get_financials()
        expense = Projector(
            financials, DCFAssumptions.from_yaml(overrides={"sbc": {"method": "expense"}})
        )
        dilute = Projector(
            financials,
            DCFAssumptions.from_yaml(
                overrides={"sbc": {"method": "dilute", "buyback_offset_pct": 1.0}}
            ),
        )
        assert list(dilute.unlevered_fcf) == pytest.approx(list(expense.unlevered_fcf))


class TestM5FootballFieldBridge:
    """The chart omitted preferred equity and investments from its own bridge."""

    def test_rows_agree_with_the_headline_when_bridge_items_are_set(self):
        warnings.simplefilter("ignore")
        financials = YFinanceClient("AAPL", offline_mode=True).get_financials()
        assumptions = DCFAssumptions.from_yaml(
            overrides={"bridge": {"preferred_equity": 50e9, "investments": 20e9}}
        )
        result = DCFEngine(financials, assumptions, ticker="AAPL").run()
        field = football_field(result)

        gordon_row = field[field["method"] == "DCF (perpetuity growth TV)"]
        assert not gordon_row.empty
        assert float(gordon_row["midpoint"].iloc[0]) == pytest.approx(
            result.value_per_share, rel=1e-9
        )

    def test_point_rows_have_a_visible_span(self):
        """A zero-length visible segment renders as nothing in a stacked bar chart."""
        warnings.simplefilter("ignore")
        financials = YFinanceClient("AAPL", offline_mode=True).get_financials()
        result = DCFEngine(financials, DCFAssumptions.from_yaml(), ticker="AAPL").run()
        field = football_field(result)
        assert (field["high"] > field["low"]).all(), "every bar needs a non-zero span"


class TestM14TerminalReinvestment:
    def test_capex_far_above_depreciation_warns(self):
        warnings.simplefilter("ignore")
        financials = YFinanceClient("MSFT", offline_mode=True).get_financials()
        result = DCFEngine(financials, DCFAssumptions.from_yaml(), ticker="MSFT").run()
        assert any("depreciation" in w for w in result.warnings), (
            "carrying a growth-phase reinvestment gap into perpetuity must be flagged"
        )


class TestM12Coercion:
    def test_currency_prefixed_accounting_negative(self):
        assert to_float("$(500)") == pytest.approx(-500.0)
        assert to_float("(£1,200)") == pytest.approx(-1200.0)
        assert to_float("$(2.5B)") == pytest.approx(-2.5e9)


class TestSbcEquivalenceStillHolds:
    """The headline claim, pinned at the true measured tolerance."""

    @pytest.mark.parametrize("ticker", ["AAPL", "MSFT", "TSLA"])
    def test_methods_agree_within_the_documented_band(self, ticker):
        warnings.simplefilter("ignore")
        financials = YFinanceClient(ticker, offline_mode=True).get_financials()
        values = {}
        for method in ("expense", "dilute"):
            assumptions = DCFAssumptions.from_yaml(overrides={"sbc": {"method": method}})
            values[method] = DCFEngine(
                financials, assumptions, ticker=ticker
            ).run().value_per_share

        gap = abs(values["dilute"] / values["expense"] - 1.0)
        # 5% binds: the measured worst case is TSLA at 3.9%. The old 15% threshold
        # passed even with dilution removed entirely.
        assert gap < 0.05, f"{ticker} methods diverged by {gap:.2%}"


def test_offline_fixtures_still_present():
    """Guards the offline promise: the committed data must actually be committed."""
    for ticker in ("AAPL", "MSFT", "TSLA"):
        assert (Path("data/offline_sample") / ticker / "income_stmt.json").exists()
