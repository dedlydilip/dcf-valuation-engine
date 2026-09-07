"""Execute exported Excel formulas with the independent formulas interpreter.

This checks actual stored expressions, not duplicated Python formula strings.
Native Excel remains a separate compatibility check.
"""

import pytest
from openpyxl import load_workbook

from src.dcf.engine import DCFEngine
from src.dcf.sensitivity import wacc_vs_exit_multiple, wacc_vs_growth
from src.excel.builder import ExcelModelBuilder
from src.models.assumptions import DCFAssumptions


def calculate(path):
    formulas = pytest.importorskip("formulas")
    return formulas.ExcelModel().loads(str(path)).finish().calculate()


def read(solution, ref):
    sheet, cell = ref.replace("$", "").split("!")
    ending = f"]{sheet}'!{cell}".upper()
    matches = [v for k, v in solution.items() if str(k).upper().endswith(ending)]
    assert len(matches) == 1, (ref, list(solution))
    return matches[0].value.item()


@pytest.mark.parametrize("method", ["gordon", "value_driver", "exit_multiple"])
@pytest.mark.parametrize("sbc", ["expense", "dilute"])
def test_exported_formula_value_matches_engine(tmp_path, tech_financials, method, sbc):
    a = DCFAssumptions.model_validate(
        {
            "projection": {"starting_nol": 100, "margin_basis": "before_sbc", "ebit_margin": 0.3},
            "wacc": {
                "capital_structure": "target",
                "target_debt_weight": 0.3,
                "floor_cost_of_equity": True,
            },
            "terminal": {"method": method, "terminal_fcf_mode": "value_driver", "ronic": 0.15},
            "sbc": {"method": sbc, "buyback_offset_pct": 0.4, "option_overhang_shares": 3},
        }
    )
    result = DCFEngine(tech_financials, a).run()
    builder = ExcelModelBuilder(result, financials=tech_financials)
    path = tmp_path / "case.xlsx"
    builder.build(path)
    book = load_workbook(path, data_only=False)
    solution = calculate(path)
    for key, expected in [
        ("vps", result.value_per_share),
        ("equity", result.bridge.equity_value),
        ("wacc", result.wacc.wacc),
        ("coe", result.wacc.cost_of_equity),
        ("tv_used", result.terminal.value),
        ("shares_out", result.bridge.shares),
    ]:
        actual = read(solution, builder.refs[key])
        assert actual == pytest.approx(expected, rel=1e-10), (key, actual, expected)
    for label, function in [
        ("WACC \\ growth", wacc_vs_growth),
        ("WACC \\ multiple", wacc_vs_exit_multiple),
    ]:
        header = next(row[0].row for row in book["Sensitivity"] if row[0].value == label)
        grid = function(tech_financials, a, result).frame
        for i in range(grid.shape[0]):
            for j in range(grid.shape[1]):
                from openpyxl.utils import get_column_letter

                actual = read(solution, f"Sensitivity!{get_column_letter(j + 2)}{header + 1 + i}")
                assert actual == pytest.approx(grid.iloc[i, j], rel=1e-10), (
                    label,
                    i,
                    j,
                    actual,
                    grid.iloc[i, j],
                )


def test_golden_formula_value_is_hand_calculated(tmp_path, golden_financials, golden_assumptions):
    result = DCFEngine(golden_financials, golden_assumptions).run()
    builder = ExcelModelBuilder(result, financials=golden_financials)
    path = tmp_path / "golden.xlsx"
    builder.build(path)
    value = read(calculate(path), builder.refs["vps"])
    # Each of three explicit PVs=150; terminal 199.65*1.02/.08 /1.1^3=1912.5.
    assert value == pytest.approx((450 + 1912.5) / 100, abs=1e-10)


def test_invalid_workbook_edits_refuse_a_price(tmp_path, golden_financials, golden_assumptions):
    result = DCFEngine(golden_financials, golden_assumptions).run()
    builder = ExcelModelBuilder(result, financials=golden_financials)
    book = builder.workbook()
    sheet, cell = builder.refs["buyback"].replace("$", "").split("!")
    book[sheet][cell] = -0.5
    path = tmp_path / "invalid.xlsx"
    book.save(path)
    solution = calculate(path)
    assert str(read(solution, builder.refs["vps"])) == "#N/A"


@pytest.mark.parametrize(
    "ref,path,value",
    [
        ("buyback", "sbc.buyback_offset_pct", 0.8),
        ("ronic", "terminal.ronic", 0.2),
        ("sbc_method", "sbc.method", "expense"),
        ("tv_mode", "terminal.terminal_fcf_mode", "fcf5"),
        ("floor_mult", "terminal.mature_industry_multiple", 4),
        ("decay", "terminal.decay_turns_per_pp", 0.3),
        ("target_wd", "wacc.target_debt_weight", 0.4),
        ("coe_floor", "wacc.floor_cost_of_equity", True),
    ],
)
def test_live_input_edit_matches_forward_engine(tmp_path, tech_financials, ref, path, value):
    a = DCFAssumptions.model_validate(
        {
            "projection": {"ebit_margin": 0.3, "margin_basis": "before_sbc"},
            "wacc": {
                "capital_structure": "target",
                "beta_override": -0.3,
                "discount_rate_override": 0.08,
            },
            "sbc": {"method": "dilute"},
            "terminal": {"method": "exit_multiple", "terminal_fcf_mode": "value_driver"},
        }
    )
    peers = {"ev_ebitda_median": 20, "revenue_growth_median": 0.15, "peer_count": 4, "min_peers": 3}
    if ref == "ronic":
        a.terminal.method = "value_driver"
    elif ref == "tv_mode":
        a.terminal.method = "gordon"
    base = DCFEngine(tech_financials, a, peers).run()
    builder = ExcelModelBuilder(base, financials=tech_financials)
    book = builder.workbook()
    sheet, cell = builder.refs[ref].replace("$", "").split("!")
    book[sheet][cell] = int(value) if isinstance(value, bool) else value
    file = tmp_path / "edited.xlsx"
    book.save(file)
    patched = a.model_copy(deep=True)
    section, key = path.split(".")
    setattr(getattr(patched, section), key, value)
    expected = DCFEngine(tech_financials, patched, peers).run()
    if ref in ("ronic", "tv_mode"):
        assert expected.value_per_share != pytest.approx(base.value_per_share)
    assert read(calculate(file), builder.refs["vps"]) == pytest.approx(
        expected.value_per_share, rel=1e-10
    )
