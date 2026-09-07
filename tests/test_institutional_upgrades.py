from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.dcf.engine import DCFEngine
from src.dcf.projector import Projector
from src.dcf.terminal_value import TerminalValue
from src.excel.builder import ExcelModelBuilder
from src.fetcher.yfinance_client import YFinanceClient
from src.models.assumptions import DCFAssumptions, TerminalAssumptions


def test_mckinsey_value_driver_steady_state_identity():
    """Verify that when RONIC == WACC, TV equals NOPAT_{N+1} / WACC."""
    assump = TerminalAssumptions(perpetuity_growth=0.03, ronic=None)
    calc = TerminalValue(assump)

    terminal_nopat = 100.0
    wacc = 0.10
    g = 0.03
    nopat_next = terminal_nopat * (1.0 + g)  # 103.0

    # Under RONIC == WACC:
    # Reinvestment rate = g / WACC = 0.03 / 0.10 = 0.30
    # FCF_steady_state = 103.0 * (1 - 0.30) = 72.10
    # TV = 72.10 / (0.10 - 0.03) = 72.10 / 0.07 = 1030.0
    # Exactly equal to NOPAT_{N+1} / WACC = 103.0 / 0.10 = 1030.0!
    res = calc.value_driver_value(terminal_nopat, wacc=wacc, growth=g, ronic=wacc)
    assert res.ok
    expected_tv = nopat_next / wacc
    assert pytest.approx(res.value, rel=1e-6) == expected_tv
    assert res.decay_detail["reinvestment_rate"] == pytest.approx(0.30)


def test_mckinsey_value_driver_moat_vs_destruction():
    """Verify that RONIC > WACC creates value with growth, while RONIC < WACC destroys value."""
    assump = TerminalAssumptions(perpetuity_growth=0.025)
    calc = TerminalValue(assump)

    terminal_nopat = 100.0
    wacc = 0.10

    # Wide Moat: RONIC = 20% > WACC (10%)
    tv_moat_low_g = calc.value_driver_value(
        terminal_nopat, wacc=wacc, growth=0.02, ronic=0.20
    ).value
    tv_moat_high_g = calc.value_driver_value(
        terminal_nopat, wacc=wacc, growth=0.04, ronic=0.20
    ).value
    # Higher growth with positive spread must CREATE value:
    assert tv_moat_high_g > tv_moat_low_g

    # Value Destructive: RONIC = 5% < WACC (10%)
    tv_destr_low_g = calc.value_driver_value(
        terminal_nopat, wacc=wacc, growth=0.02, ronic=0.05
    ).value
    tv_destr_high_g = calc.value_driver_value(
        terminal_nopat, wacc=wacc, growth=0.04, ronic=0.05
    ).value
    # Higher growth with negative spread must DESTROY value:
    assert tv_destr_high_g < tv_destr_low_g


def test_capex_fade_convergence():
    """Verify that fade_capex_to_da smoothly converges CapEx to terminal_capex_to_da * D&A."""
    client = YFinanceClient("MSFT", offline_mode=True)
    fin = client.get_financials()

    # Standard model without fade
    base_assump = DCFAssumptions.from_yaml("config/assumptions.yaml", scenario="base")
    p_standard = Projector(fin, base_assump).result

    # Model with fade enabled: target 1.0x D&A in Year 5
    fade_assump = DCFAssumptions.from_yaml(
        "config/assumptions.yaml",
        scenario="base",
        overrides={"projection": {"fade_capex_to_da": True, "terminal_capex_to_da": 1.0}},
    )
    p_fade = Projector(fin, fade_assump).result

    capex_std_yr5 = float(p_standard.table.loc["capex"].iloc[-1])
    capex_fade_yr5 = float(p_fade.table.loc["capex"].iloc[-1])
    da_fade_yr5 = float(p_fade.table.loc["da"].iloc[-1])

    # In Microsoft's standard model, CapEx in yr 5 is ~2.5x D&A
    assert capex_std_yr5 > 2.0 * da_fade_yr5

    # With fade enabled, Year 5 CapEx must equal exactly 1.0 * Year 5 D&A
    assert pytest.approx(capex_fade_yr5, rel=1e-3) == da_fade_yr5
    # And FCF in Year 5 must be substantially higher under normalized capex
    assert float(p_fade.unlevered_fcf.iloc[-1]) > float(p_standard.unlevered_fcf.iloc[-1])


def test_include_investments_bridge():
    """Verify that include_investments adds marketable securities to equity value."""
    client = YFinanceClient("AAPL", offline_mode=True)
    fin = client.get_financials()

    # Default: excluded
    assump_excl = DCFAssumptions.from_yaml("config/assumptions.yaml", scenario="base")
    res_excl = DCFEngine(fin, assump_excl, ticker="AAPL").run()
    assert res_excl.bridge.investments == 0.0

    # With include_investments: True
    assump_incl = DCFAssumptions.from_yaml(
        "config/assumptions.yaml",
        scenario="base",
        overrides={"bridge": {"include_investments": True}},
    )
    res_incl = DCFEngine(fin, assump_incl, ticker="AAPL").run()

    # Apple reports $77,723,000,000 in long-term investments
    assert res_incl.bridge.investments == 77_723_000_000.0
    # Equity value must be higher by exactly the portfolio value
    assert (
        pytest.approx(res_incl.bridge.equity_value - res_excl.bridge.equity_value)
        == 77_723_000_000.0
    )
    # Value per share increases from ~$120.08 to ~$125.41 (+4.4%)
    assert res_incl.value_per_share > res_excl.value_per_share
    pct_gain = (res_incl.value_per_share / res_excl.value_per_share) - 1.0
    assert 0.04 < pct_gain < 0.05


def test_excel_summary_sheet_parity(tmp_path: Path):
    """Verify that the generated Excel Summary sheet contains Moat, Expectations, and Margin of Safety."""
    client = YFinanceClient("AAPL", offline_mode=True)
    fin = client.get_financials()
    assump = DCFAssumptions.from_yaml("config/assumptions.yaml", scenario="base")
    res = DCFEngine(fin, assump, ticker="AAPL").run()

    builder = ExcelModelBuilder(res, financials=fin)
    xlsx_path = tmp_path / "AAPL_test.xlsx"
    builder.build(xlsx_path)

    wb = load_workbook(xlsx_path, data_only=False)
    assert "Summary" in wb.sheetnames
    ws = wb["Summary"]

    found_moat = False
    found_expectations = False
    found_target_entry = False

    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is not None:
                text = str(cell)
                if "Accounting Capital Efficiency (Moat Unverified)" in text:
                    found_moat = True
                if "Market Expectations & Margin of Safety" in text:
                    found_expectations = True
                if "Illustrative 15% discount to modeled value" in text:
                    found_target_entry = True

    assert found_moat, "Summary sheet missing Economic Moat section"
    assert found_expectations, "Summary sheet missing Market Expectations section"
    assert found_target_entry, "Summary sheet missing Target Entry margin of safety section"
