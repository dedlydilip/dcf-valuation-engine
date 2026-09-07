"""Export live Excel formulas for the core valuation.

Forecasts, cash taxes, discount rates, terminal selection and dilution follow
the Python accounting policy. Supporting source observations and analytical
reports are snapshots at generation and are explicitly labeled for rebuilding.
"""

from __future__ import annotations

import math
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.excel import styles as S

FIRST_FORECAST_COL = 4  # column D; column C holds the base year


class SheetCursor:
    """Row-by-row writer that remembers where it put things."""

    def __init__(self, worksheet: Worksheet, refs: dict[str, str]) -> None:
        self.ws = worksheet
        self.refs = refs
        self.row = 1

    # ------------------------------------------------------------------ helpers

    def addr(self, row: int, col: int) -> str:
        return f"{self.ws.title}!${get_column_letter(col)}${row}"

    def cell_addr(self, row: int, col: int) -> str:
        return f"${get_column_letter(col)}${row}"

    def remember(self, key: str, row: int, col: int) -> str:
        ref = self.addr(row, col)
        self.refs[key] = ref
        return ref

    def skip(self, n: int = 1) -> None:
        self.row += n

    # ------------------------------------------------------------------- blocks

    def title(self, text: str, width: int = 9) -> None:
        cell = self.ws.cell(row=self.row, column=1, value=text)
        cell.font = S.TITLE_FONT
        for col in range(1, width + 1):
            self.ws.cell(row=self.row, column=col).fill = S.TITLE_FILL
        self.ws.row_dimensions[self.row].height = 22
        self.row += 2

    def section(self, text: str, width: int = 9) -> None:
        cell = self.ws.cell(row=self.row, column=1, value=text)
        cell.font = S.SECTION_FONT
        for col in range(1, width + 1):
            self.ws.cell(row=self.row, column=col).fill = S.SECTION_FILL
        self.row += 1

    def note(self, text: str, warn: bool = False) -> None:
        from openpyxl.styles import Alignment

        cell = self.ws.cell(row=self.row, column=1, value=text)
        cell.font = S.WARN_FONT if warn else S.NOTE_FONT
        width = max(4, self.ws.max_column)
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=width)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        chars = sum(
            self.ws.column_dimensions[get_column_letter(i)].width or 13 for i in range(1, width + 1)
        )
        self.ws.row_dimensions[self.row].height = max(
            28, min(400, math.ceil(len(text) / max(chars, 40)) * 15)
        )
        self.row += 1

    def label_value(
        self,
        label: str,
        value: Any,
        fmt: str | None = None,
        key: str | None = None,
        kind: str = "input",
        col: int = 2,
    ) -> str:
        self.ws.cell(row=self.row, column=1, value=label).font = S.LABEL_FONT
        cell = self.ws.cell(row=self.row, column=col, value=_num(value))
        cell.font = _font_for(kind)
        if fmt:
            cell.number_format = fmt
        ref = self.remember(key, self.row, col) if key else self.addr(self.row, col)
        self.row += 1
        return ref

    def series_row(
        self,
        label: str,
        values: list[Any],
        fmt: str | None = None,
        key: str | None = None,
        kind: str = "formula",
        start_col: int = FIRST_FORECAST_COL,
        bold: bool = False,
        border: Any = None,
    ) -> int:
        """Write a labelled row of per-year cells and return its row number."""
        label_cell = self.ws.cell(row=self.row, column=1, value=label)
        label_cell.font = S.TOTAL_FONT if bold else S.LABEL_FONT
        for offset, value in enumerate(values):
            cell = self.ws.cell(row=self.row, column=start_col + offset, value=_num(value))
            cell.font = S.TOTAL_FONT if bold else _font_for(kind)
            if fmt:
                cell.number_format = fmt
            if border is not None:
                cell.border = border
        row = self.row
        if key:
            self.refs[key] = f"{self.ws.title}!${get_column_letter(start_col)}${row}"
            self.refs[f"{key}_row"] = str(row)
            self.refs[f"{key}_start_col"] = str(start_col)
            self.refs[f"{key}_end_col"] = str(start_col + len(values) - 1)
        self.row += 1
        return row

    def headers(self, labels: list[str], start_col: int = 1) -> None:
        for offset, text in enumerate(labels):
            cell = self.ws.cell(row=self.row, column=start_col + offset, value=text)
            cell.font = S.HEADER_FONT
            cell.fill = S.HEADER_FILL
            cell.alignment = S.CENTER
        self.row += 1


def _font_for(kind: str):
    return {
        "input": S.INPUT_FONT,
        "formula": S.FORMULA_FONT,
        "link": S.LINK_FONT,
        "total": S.TOTAL_FONT,
    }.get(kind, S.FORMULA_FONT)


class ExcelModelBuilder:
    """Assemble the workbook from a completed valuation."""

    def __init__(
        self,
        result: Any,
        financials: Any = None,
        comps: Any = None,
        monte_carlo: Any = None,
        football: pd.DataFrame | None = None,
    ) -> None:
        self.result = result
        self.financials = financials
        self.comps = comps
        self.monte_carlo = monte_carlo
        self.football = football
        self.refs: dict[str, str] = {}
        self.years = len(result.pv_explicit)
        currency = getattr(financials, "currency", "USD")
        self.price_format = f'"{currency} "#,##0.00;("{currency} "#,##0.00)'

    def _base_shares(self) -> float:
        """Share count before any forecast issuance.

        The workbook applies dilution itself, so it needs the undiluted starting
        point rather than the engine's already-diluted output.
        """
        from src.dcf.bridge import base_share_count
        from src.models.errors import ValuationError

        if self.financials is not None:
            try:
                return base_share_count(self.financials)
            except (ValueError, ValuationError):
                pass
        if self.result.assumptions.sbc.grow_share_count:
            raise ValueError(
                "financials are required to build a dilute-method workbook, because the "
                "undiluted share count cannot be recovered from the diluted result"
            )
        overhang = self.result.assumptions.sbc.option_overhang_shares or 0.0
        return self.result.bridge.shares - overhang

    # -------------------------------------------------------------------- build

    def build(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        book = self.workbook()
        book.save(path)
        import json

        from src.models.provenance import json_safe

        path.with_suffix(".manifest.json").write_text(
            json.dumps(json_safe(self.result.manifest), indent=2, allow_nan=False), encoding="utf-8"
        )
        return path

    def workbook(self) -> Workbook:
        """Build the workbook in memory, without writing it.

        Separate from `build` so the non-finite sweep below -- and the tests that
        check it -- can run against a workbook that still holds its float values.
        Once openpyxl has written and reloaded a file, every non-finite number has
        become None, so nothing downstream of `save` can tell whether one was there.
        """
        book = Workbook()
        book.remove(book.active)

        # Order matters: sheets that are referenced must exist before the formulas
        # that point at them are written.
        self._inputs(book.create_sheet("Inputs"))
        self._wacc(book.create_sheet("WACC"))
        self._dcf(book.create_sheet("DCF"))
        self._sensitivity(book.create_sheet("Sensitivity"))
        self._comps(book.create_sheet("Comps"))
        self._football(book.create_sheet("FootballField"))
        self._historicals(book.create_sheet("Historicals"))
        self._summary(book.create_sheet("Summary", 0))

        import json

        from openpyxl.styles import Alignment

        provenance = book.create_sheet("Provenance")
        provenance.append(["Run manifest at generation; regenerate after changing inputs"])
        for key, value in self.result.manifest.items():
            provenance.append(
                [
                    key,
                    json.dumps(value, sort_keys=True, default=str)
                    if isinstance(value, (dict, list))
                    else str(value),
                ]
            )
        provenance.column_dimensions["A"].width = 28
        provenance.column_dimensions["B"].width = 100
        for row in provenance:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        assert_all_cells_finite(book)
        return book

    # ------------------------------------------------------------------- inputs

    def _inputs(self, ws: Worksheet) -> None:
        S.col_width(ws, {"A": 42, "B": 16, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14})
        cur = SheetCursor(ws, self.refs)
        result = self.result
        assumptions = result.assumptions
        proj = result.projection
        table = proj.table
        info = getattr(self.financials, "info", None) or {}

        cur.title(f"{result.ticker} - Valuation Inputs", width=8)
        cur.note(
            f"Amounts in {getattr(self.financials, 'currency', 'USD')} millions; shares in millions; prices per share. Core valuation updates with blue inputs. Supporting snapshot reports require rebuilding."
        )
        cur.skip()

        cur.section("Company", width=8)
        cur.label_value("Ticker", result.ticker, key="ticker")
        cur.label_value("Name", info.get("longName", result.ticker))
        cur.label_value("Scenario", assumptions.scenario)
        cur.label_value("Data source", getattr(self.financials, "source", "n/a"))
        # A valuation with no date on it is unreadable six months later: the reader
        # cannot tell whether a stale-looking price is an error or simply old. These
        # are two different dates and both matter -- when the model was run, and how
        # old the statements behind it are.
        cur.label_value("Workbook generated on", datetime.now().strftime("%Y-%m-%d"))
        cur.label_value("Latest fiscal period end", _latest_period(self.financials))
        cur.label_value(
            "Snapshot share price",
            result.bridge.current_price or 0.0,
            self.price_format,
            key="price",
        )
        cur.skip()

        cur.section("Discount rate build-up", width=8)
        w = result.wacc
        cur.label_value("Risk-free rate", w.risk_free_rate, S.PERCENT_2, key="rf")
        cur.label_value("Equity risk premium", w.equity_risk_premium, S.PERCENT_2, key="erp")
        cur.label_value(
            "Raw levered beta at current capital structure", w.raw_beta, S.RATIO, key="beta"
        )
        cur.label_value("Size premium", w.size_premium, S.PERCENT_2, key="size_prem")
        cur.label_value("Country risk premium", w.country_risk_premium, S.PERCENT_2, key="crp")
        cur.label_value("Cost of debt (pre-tax)", w.cost_of_debt, S.PERCENT_2, key="kd")
        cur.label_value("Tax rate", assumptions.projection.tax_rate, S.PERCENT, key="tax")
        cur.label_value("Market capitalisation", w.market_cap, S.MONEY_MM, key="mcap")
        cur.label_value("Total debt", w.total_debt, S.MONEY_MM, key="debt")
        cur.label_value("Cash and short-term investments", w.cash, S.MONEY_MM, key="cash")
        cur.label_value(
            "Capital structure (current / target)",
            assumptions.wacc.capital_structure,
            key="cap_structure",
        )
        cur.label_value(
            "Target debt weight (used only when target)",
            assumptions.wacc.target_debt_weight,
            S.PERCENT,
            key="target_wd",
        )
        cur.skip()

        cur.section("Share count", width=8)
        # Must be the UNDILUTED count. `result.bridge.shares` is already diluted when
        # the run used sbc.method = "dilute", and feeding that back into the
        # workbook's own dilution formula would dilute it a second time -- the same
        # double-count, one level up.
        cur.label_value(
            "Diluted shares (current, before forecast issuance)",
            self._base_shares(),
            S.SHARES_MM,
            key="shares_in",
        )
        cur.label_value(
            "Option overhang (treasury method)",
            assumptions.sbc.option_overhang_shares or 0.0,
            S.SHARES_MM,
            key="overhang",
        )
        cur.skip()

        cur.section("Stock-based compensation policy", width=8)
        cur.label_value("SBC method (expense / dilute)", assumptions.sbc.method, key="sbc_method")
        cur.note(
            "expense: SBC stays a cost inside EBIT, share count held flat. "
            "dilute: SBC added back, share count grows. Never both -- that double-counts."
        )
        cur.skip()

        cur.section("Projection drivers", width=8)
        year_labels = ["Base"] + [f"Year {i}" for i in range(1, self.years + 1)]
        cur.headers(year_labels, start_col=3)

        base_revenue = proj.drivers["base_revenue"]
        cur.series_row(
            "Revenue growth",
            [""] + [float(v) for v in table.loc["revenue_growth"]],
            S.PERCENT,
            key="growth",
            kind="input",
            start_col=3,
        )
        cur.series_row(
            "EBIT margin (basis specified below)",
            [""]
            + [
                float(e / r)
                for e, r in zip(
                    table.loc["ebit_pre_sbc"]
                    if assumptions.projection.margin_basis == "before_sbc"
                    else table.loc["ebit"],
                    table.loc["revenue"],
                    strict=True,
                )
            ],
            S.PERCENT,
            key="margin",
            kind="input",
            start_col=3,
        )
        cur.series_row(
            "D&A % of revenue",
            [""]
            + [
                float(v) / float(r)
                for v, r in zip(table.loc["da"], table.loc["revenue"], strict=True)
            ],
            S.PERCENT,
            key="da_pct",
            kind="input",
            start_col=3,
        )
        cur.series_row(
            "Capex % of revenue",
            [""]
            + [
                float(v) / float(r)
                for v, r in zip(table.loc["capex"], table.loc["revenue"], strict=True)
            ],
            S.PERCENT,
            key="capex_pct",
            kind="input",
            start_col=3,
        )
        cur.series_row(
            "Net working capital % of revenue",
            [proj.drivers["base_nwc"] / base_revenue]
            + [
                float(v) / float(r)
                for v, r in zip(table.loc["nwc"], table.loc["revenue"], strict=True)
            ],
            S.PERCENT,
            key="nwc_pct",
            kind="input",
            start_col=3,
        )
        cur.series_row(
            "SBC % of revenue",
            [""]
            + [
                float(v) / float(r)
                for v, r in zip(table.loc["sbc"], table.loc["revenue"], strict=True)
            ],
            S.PERCENT,
            key="sbc_pct",
            kind="input",
            start_col=3,
        )
        cur.label_value("Base year revenue", base_revenue, S.MONEY_MM, key="base_revenue")
        cur.skip()

        cur.section("Terminal value", width=8)
        term = assumptions.terminal
        cur.label_value("Perpetuity growth rate", term.perpetuity_growth, S.PERCENT_2, key="g")
        exit_result = result.terminal_all.get("exit_multiple")
        cur.label_value(
            "Static exit multiple assumption",
            term.static_exit_multiple,
            S.MULTIPLE,
            key="static_mult",
        )
        cur.label_value(
            "Mature industry floor", term.mature_industry_multiple, S.MULTIPLE, key="floor_mult"
        )
        cur.label_value(
            "Decay (turns per pp of growth lost)", term.decay_turns_per_pp, S.RATIO, key="decay"
        )
        cur.label_value("Terminal method preference", term.method, key="tv_pref")
        cur.label_value(
            "Mid-year convention (1 = yes)",
            1 if assumptions.projection.mid_year_convention else 0,
            S.INTEGER,
            key="midyear",
        )
        cur.skip()

        cur.section("Bridge items", width=8)
        cur.label_value(
            "Minority interest", result.bridge.minority_interest, S.MONEY_MM, key="minority"
        )
        cur.label_value(
            "Preferred equity", result.bridge.preferred_equity, S.MONEY_MM, key="preferred"
        )
        cur.label_value(
            "Non-operating investments", result.bridge.investments, S.MONEY_MM, key="invest"
        )

        cur.skip()
        cur.section("Calculation policies", width=8)
        cur.label_value("Terminal FCF mode", term.terminal_fcf_mode, key="tv_mode")
        cur.label_value("RONIC (0 = WACC)", term.ronic or 0, S.PERCENT_2, key="ronic")
        cur.label_value(
            "Buyback offset fraction", assumptions.sbc.buyback_offset_pct, S.PERCENT, key="buyback"
        )
        cur.label_value(
            "Floor cost of equity at risk-free (1=yes)",
            int(assumptions.wacc.floor_cost_of_equity),
            key="coe_floor",
        )
        cur.label_value(
            "Direct WACC override (0 = calculate)",
            assumptions.wacc.discount_rate_override or 0,
            S.PERCENT_2,
            key="wacc_override",
        )
        cur.label_value(
            "Opening tax loss carryforward",
            assumptions.projection.starting_nol,
            S.MONEY_MM,
            key="nol",
        )
        cur.label_value(
            "Operating margin basis", assumptions.projection.margin_basis, key="margin_basis"
        )
        cur.label_value(
            "Historical opening operating NWC",
            f"={self.refs['base_revenue']}*{self.refs['nwc_pct']}",
            S.MONEY_MM,
            key="base_nwc",
            kind="formula",
        )
        r = self.refs
        detail = exit_result.decay_detail if exit_result else {}
        cur.label_value("Exit multiple mode", term.exit_multiple_mode, key="exit_mode")
        cur.label_value(
            "Peer anchor supported (1=yes)",
            int(detail.get("anchor_is_peer_median", 0)),
            key="peer_supported",
        )
        cur.label_value(
            "Screened peer EV/EBITDA anchor",
            detail.get("peer_median_multiple", term.static_exit_multiple),
            S.MULTIPLE,
            key="peer_mult",
        )
        cur.label_value(
            "Screened peer revenue growth",
            detail.get("peer_median_growth", 0),
            S.PERCENT_2,
            key="peer_growth",
        )
        last_growth = _cell(r["growth"], self.years)
        anchor = f"IF({r['peer_supported']}=1,{r['peer_mult']},{r['static_mult']})"
        peer_growth = f"IF({r['peer_supported']}=1,{r['peer_growth']},{last_growth})"
        multiple_formula = f'=IF({r["exit_mode"]}="static",{r["static_mult"]},MIN(MAX({anchor}-{r["decay"]}*MAX(0,({peer_growth}-{last_growth})*100),MIN({r["floor_mult"]},{anchor})),{anchor}))'
        cur.label_value(
            "Exit multiple used (EV/EBITDA)",
            multiple_formula,
            S.MULTIPLE,
            key="exit_mult",
            kind="formula",
        )
        cur.note(
            "Peer observations are fixed at generation. Changing growth, decay, floor or mode updates the resolved multiple; refresh source data by rebuilding."
        )
        cur.note(
            "NOL schedule assumes unrestricted carryforward without expiry or utilization caps. Cash repurchases receive no tax deduction."
        )
        checks = [
            f"{r['base_revenue']}>0",
            f"{r['shares_in']}>0",
            f"{r['overhang']}>=0",
            f'OR({r["sbc_method"]}="expense",{r["sbc_method"]}="dilute")',
            f'OR({r["margin_basis"]}="after_sbc",{r["margin_basis"]}="before_sbc")',
            f'OR({r["cap_structure"]}="current",{r["cap_structure"]}="target")',
            f'OR({r["tv_mode"]}="fcf5",{r["tv_mode"]}="value_driver")',
            f'OR({r["tv_pref"]}="gordon",{r["tv_pref"]}="both",{r["tv_pref"]}="value_driver",{r["tv_pref"]}="exit_multiple")',
            f'OR({r["exit_mode"]}="static",{r["exit_mode"]}="dynamic")',
        ]
        for key, lo, hi in [
            ("buyback", 0, 1),
            ("tax", 0, 1),
            ("rf", 0, 0.25),
            ("erp", 0, 0.2),
            ("beta", -5, 10),
            ("target_wd", 0, 0.95),
            ("g", -0.02, 0.06),
            ("ronic", 0, 2),
            ("static_mult", 0, 100),
            ("floor_mult", 0, 50),
            ("decay", 0, 20),
            ("wacc_override", 0, 1),
            ("kd", 0, 0.5),
            ("size_prem", 0, 0.1),
            ("crp", 0, 0.2),
            ("nol", 0, 1e300),
            ("debt", 0, 1e300),
            ("cash", 0, 1e300),
        ]:
            lower = ">" if key in ("erp", "static_mult", "floor_mult") else ">="
            upper = "<" if key == "tax" else "<="
            checks.extend([f"{r[key]}{lower}{lo}", f"{r[key]}{upper}{hi}"])
        for key, lo, hi in [
            ("growth", -1, 10),
            ("margin", -5, 1),
            ("da_pct", 0, 5),
            ("capex_pct", 0, 5),
            ("nwc_pct", -5, 5),
            ("sbc_pct", 0, 1e300),
        ]:
            cells = f"{_cell(r[key], 1)}:{_cell(r[key], self.years).split('!')[1]}"
            lower = ">" if key == "growth" else ">="
            checks.extend(
                [f"COUNT({cells})={self.years}", f"MIN({cells}){lower}{lo}", f"MAX({cells})<={hi}"]
            )
        cur.label_value(
            "Input policy valid (FALSE blocks valuation)",
            "=AND(" + ",".join(checks) + ")",
            key="inputs_valid",
            kind="formula",
        )
        ws.freeze_panes = "A3"

    # --------------------------------------------------------------------- wacc

    def _wacc(self, ws: Worksheet) -> None:
        S.col_width(ws, {"A": 46, "B": 20})
        cur = SheetCursor(ws, self.refs)
        r = self.refs
        cur.title("Weighted Average Cost of Capital", width=4)
        cur.note(
            "Beta adjusts with target leverage. Debt cost is an analyst input; historical book yield is a proxy, not a current market yield."
        )
        cur.label_value("Risk-free rate", f"={r['rf']}", S.PERCENT_2, kind="link")
        beta = cur.label_value(
            "Beta",
            f'=IF({r["cap_structure"]}="target",{r["beta"]}/(1+(1-{r["tax"]})*{r["debt"]}/{r["mcap"]})*(1+(1-{r["tax"]})*{r["target_wd"]}/(1-{r["target_wd"]})),{r["beta"]})',
            S.RATIO,
            key="levered_beta",
        )
        raw = f"{r['rf']}+{beta}*{r['erp']}+{r['size_prem']}+{r['crp']}"
        cur.label_value(
            "Cost of equity",
            f"=IF({r['coe_floor']}=1,MAX({r['rf']},{raw}),{raw})",
            S.PERCENT_2,
            key="coe",
        )
        cur.label_value("Pre-tax cost of debt", f"={r['kd']}", S.PERCENT_2, kind="link")
        cur.label_value(
            "After-tax cost of debt", f"={r['kd']}*(1-{r['tax']})", S.PERCENT_2, key="atkd"
        )
        cur.label_value("Total capital", f"={r['mcap']}+{r['debt']}", S.MONEY_MM, key="totcap")
        cur.label_value(
            "Debt weight",
            f'=IF({r["cap_structure"]}="target",{r["target_wd"]},IF({r["totcap"]}>0,{r["debt"]}/{r["totcap"]},0))',
            S.PERCENT,
            key="wd",
        )
        cur.label_value("Equity weight", f"=1-{r['wd']}", S.PERCENT, key="we")
        cur.label_value(
            "WACC",
            f"=IF({r['wacc_override']}>0,{r['wacc_override']},{r['we']}*{r['coe']}+{r['wd']}*{r['atkd']})",
            S.PERCENT_2,
            key="wacc",
            kind="total",
        )

    def _dcf(self, ws: Worksheet) -> None:
        S.col_width(
            ws,
            {
                "A": 48,
                "B": 20,
                "C": 16,
                **{get_column_letter(i): 16 for i in range(4, 4 + self.years)},
            },
        )
        cur = SheetCursor(ws, self.refs)
        r = self.refs
        first = FIRST_FORECAST_COL
        last = first + self.years - 1
        letters = [get_column_letter(i) for i in range(first, last + 1)]
        cur.title(f"{self.result.ticker} - Discounted Cash Flow", width=last)
        cur.note(
            "Annual snapshot model. Blue cells are editable. Summary diagnostics are saved-at-build and require regeneration after edits."
        )
        cur.headers(["Base"] + [f"Year {i}" for i in range(1, self.years + 1)], start_col=3)

        def series(label, formulas, base=None):
            row = cur.row
            ws.cell(row, 1, label).font = S.LABEL_FONT
            if base is not None:
                ws.cell(row, 3, base).number_format = S.MONEY_MM
            for col, formula in zip(range(first, last + 1), formulas, strict=True):
                c = ws.cell(row, col, formula)
                c.font = S.FORMULA_FONT
                c.number_format = S.MONEY_MM
            cur.row += 1
            return row

        rev = cur.row
        series(
            "Revenue",
            [
                f"={('C' if i == 0 else letters[i - 1])}{rev}*(1+{_cell(r['growth'], i + 1)})"
                for i in range(self.years)
            ],
            f"={r['base_revenue']}",
        )
        sbc = series(
            "Stock-based compensation",
            [f"={letter}{rev}*{_cell(r['sbc_pct'], i + 1)}" for i, letter in enumerate(letters)],
        )
        ebit = series(
            "EBIT (GAAP, after SBC)",
            [
                f'={letter}{rev}*{_cell(r["margin"], i + 1)}-IF({r["margin_basis"]}="before_sbc",{letter}{sbc},0)'
                for i, letter in enumerate(letters)
            ],
        )
        series("EBIT before SBC", [f"={letter}{ebit}+{letter}{sbc}" for letter in letters])
        taxable = series(
            "Taxable EBIT (SBC remains deductible)", [f"={letter}{ebit}" for letter in letters]
        )
        nol_open = cur.row
        nol_use = nol_open + 1
        nol_end = nol_open + 2
        series(
            "Opening tax loss carryforward",
            [
                f"={r['nol']}" if i == 0 else f"={letters[i - 1]}{nol_end}"
                for i in range(self.years)
            ],
        )
        series(
            "Tax loss used",
            [f"=MIN({letter}{nol_open},MAX({letter}{taxable},0))" for letter in letters],
        )
        series(
            "Closing tax loss carryforward",
            [
                f"={letter}{nol_open}-{letter}{nol_use}+MAX(-{letter}{taxable},0)"
                for letter in letters
            ],
        )
        tax = series(
            "(-) Taxes",
            [f"=-(MAX({letter}{taxable},0)-{letter}{nol_use})*{r['tax']}" for letter in letters],
        )
        nopat = series("NOPAT", [f"={letter}{ebit}+{letter}{tax}" for letter in letters])
        da = series(
            "(+) Depreciation & amortisation",
            [f"={letter}{rev}*{_cell(r['da_pct'], i + 1)}" for i, letter in enumerate(letters)],
        )
        capex = series(
            "(-) Capital expenditure",
            [f"=-{letter}{rev}*{_cell(r['capex_pct'], i + 1)}" for i, letter in enumerate(letters)],
        )
        nwc = series(
            "Net working capital",
            [f"={letter}{rev}*{_cell(r['nwc_pct'], i + 1)}" for i, letter in enumerate(letters)],
            f"={r['base_nwc']}",
        )
        dnwc = series(
            "(-) Increase in working capital",
            [
                f"=-({letter}{nwc}-{('C' if i == 0 else letters[i - 1])}{nwc})"
                for i, letter in enumerate(letters)
            ],
        )
        adj = series(
            "Memo: adjusted FCF (SBC expensed)",
            [f"={letter}{nopat}+{letter}{da}+{letter}{capex}+{letter}{dnwc}" for letter in letters],
        )
        neutral = series(
            "Memo: SBC-neutral FCF (SBC added back)",
            [
                f'={letter}{adj}+{letter}{sbc}*(1-IF({r["sbc_method"]}="dilute",{r["buyback"]},0))'
                for letter in letters
            ],
        )
        core = series(
            "Unlevered FCF (per SBC method)",
            [
                f'=IF({r["sbc_method"]}="dilute",{letter}{neutral},{letter}{adj})'
                for letter in letters
            ],
        )
        idx = series("Year index", [f"={i + 1}" for i in range(self.years)])
        exp = series(
            "Discount period (years)",
            [f"={letter}{idx}-IF({r['midyear']}=1,0.5,0)" for letter in letters],
        )
        df = series("Discount factor", [f"=1/(1+{r['wacc']})^{letter}{exp}" for letter in letters])
        for row, number_format in [(idx, S.INTEGER), (exp, S.RATIO), (df, S.RATIO)]:
            for column in range(first, last + 1):
                ws.cell(row, column).number_format = number_format
        pv = series("PV of unlevered FCF", [f"={letter}{core}*{letter}{df}" for letter in letters])
        cur.label_value(
            "Sum of PV, explicit period",
            f"=SUM({letters[0]}{pv}:{letters[-1]}{pv})",
            S.MONEY_MM,
            key="sum_pv",
        )
        cur.section("Terminal value", width=last)
        letter = letters[-1]
        cur.label_value(
            "Terminal year EBITDA", f"={letter}{ebit}+{letter}{da}", S.MONEY_MM, key="term_ebitda"
        )
        cur.label_value(
            "Terminal NOPAT (SBC expensed, no perpetual NOL)",
            f"={letter}{ebit}-MAX({letter}{ebit},0)*{r['tax']}",
            S.MONEY_MM,
            key="term_nopat",
        )
        cur.label_value(
            "Terminal year FCF (SBC expensed)",
            f"={r['term_nopat']}+{letter}{da}+{letter}{capex}+{letter}{dnwc}",
            S.MONEY_MM,
            key="term_fcf",
        )
        cur.label_value(
            "Terminal value - FCF5 Gordon",
            f"=IF({r['wacc']}<={r['g']},NA(),{r['term_fcf']}*(1+{r['g']})/({r['wacc']}-{r['g']}))",
            S.MONEY_MM,
            key="tv_fcf5",
        )
        cur.label_value(
            "Terminal value - value driver",
            f"=IF(OR({r['wacc']}<={r['g']},{r['term_nopat']}<=0),NA(),{r['term_nopat']}*(1+{r['g']})*(1-{r['g']}/IF({r['ronic']}>0,{r['ronic']},{r['wacc']}))/({r['wacc']}-{r['g']}))",
            S.MONEY_MM,
            key="tv_vd",
        )
        cur.label_value(
            "Terminal value - perpetuity growth",
            f'=IF(AND({r["tv_mode"]}="value_driver",IFERROR({r["tv_vd"]}>0,FALSE)),{r["tv_vd"]},{r["tv_fcf5"]})',
            S.MONEY_MM,
            key="tv_gordon",
        )
        cur.label_value(
            "Terminal value - exit multiple",
            f"=IF({r['term_ebitda']}<=0,NA(),{r['term_ebitda']}*{r['exit_mult']})",
            S.MONEY_MM,
            key="tv_exit",
        )

        def positive(key):
            return f"IFERROR({r[key]}>0,FALSE)"

        # Same positive-value preference, then finite Gordon distress diagnostic.
        normal = f'IF({positive("tv_gordon")},"gordon",IF({positive("tv_exit")},"exit_multiple",IF({positive("tv_vd")},"value_driver","gordon")))'
        exit_first = f'IF({positive("tv_exit")},"exit_multiple",{normal})'
        vd_first = f'IF({positive("tv_vd")},"value_driver",{normal})'
        cur.label_value(
            "Terminal method actually used",
            f'=IF({r["tv_pref"]}="value_driver",{vd_first},IF({r["tv_pref"]}="exit_multiple",{exit_first},{normal}))',
            key="tv_actual",
        )
        cur.label_value(
            "Terminal value used",
            f'=IF({r["tv_actual"]}="value_driver",{r["tv_vd"]},IF({r["tv_actual"]}="exit_multiple",{r["tv_exit"]},{r["tv_gordon"]}))',
            S.MONEY_MM,
            key="tv_used",
        )
        cur.label_value(
            "PV of terminal value",
            f'={r["tv_used"]}/(1+{r["wacc"]})^IF({r["tv_actual"]}="exit_multiple",{letter}{idx},{letter}{exp})',
            S.MONEY_MM,
            key="pv_tv",
        )
        cur.label_value(
            "Implied exit multiple from perpetuity method",
            f"=IF({r['term_ebitda']}<=0,NA(),{r['tv_gordon']}/{r['term_ebitda']})",
            S.MULTIPLE,
            key="implied_mult",
        )
        cur.label_value(
            "Implied perpetuity growth from exit multiple",
            f"=IFERROR(({r['tv_exit']}*{r['wacc']}-{r['term_fcf']})/({r['tv_exit']}+{r['term_fcf']}),NA())",
            S.PERCENT_2,
            key="implied_g",
        )
        cur.section("Enterprise value to equity value", width=last)
        cur.label_value("Enterprise value", f"={r['sum_pv']}+{r['pv_tv']}", S.MONEY_MM, key="ev")
        cur.label_value(
            "Terminal value as % of EV",
            f"=IF({r['ev']}=0,NA(),{r['pv_tv']}/{r['ev']})",
            S.PERCENT,
            key="tv_pct",
        )
        cur.label_value(
            "Memo: net debt and other bridge items",
            f"={r['debt']}-{r['cash']}+{r['minority']}+{r['preferred']}-{r['invest']}",
            S.MONEY_MM,
            key="net_bridge",
        )
        cur.label_value("Equity value", f"={r['ev']}-{r['net_bridge']}", S.MONEY_MM, key="equity")
        cur.label_value(
            "Memo: PV of future SBC (at cost of equity)",
            f"=SUMPRODUCT({letters[0]}{sbc}:{letter}{sbc},1/((1+{r['coe']})^{letters[0]}{idx}:{letter}{idx}))*(1-{r['buyback']})",
            S.MONEY_MM,
            key="sbc_pv",
        )
        opening = f"({r['shares_in']}+{r['overhang']})"
        cur.label_value(
            "(/) Diluted shares",
            f'=IF({r["sbc_method"]}="dilute",IF(OR({r["equity"]}<={r["sbc_pv"]},{r["coe"]}<=-1),NA(),{opening}/(1-{r["sbc_pv"]}/{r["equity"]})),{opening})',
            S.SHARES_MM,
            key="shares_out",
        )
        cur.label_value(
            "Implied value per share",
            f"=IF(AND({r['inputs_valid']},{r['wacc']}>0,{r['wacc']}<=1),{r['equity']}/{r['shares_out']},NA())",
            self.price_format,
            key="vps",
            kind="total",
        )
        cur.label_value("Snapshot share price", f"={r['price']}", self.price_format)
        cur.label_value(
            "Upside / (downside)",
            f"=IF({r['price']}<=0,NA(),{r['vps']}/{r['price']}-1)",
            S.PERCENT,
            key="upside",
        )
        r.update(
            fcf_row=str(core),
            adj_row=str(adj),
            exp_row=str(exp),
            idx_row=str(idx),
            first_letter=letters[0],
            last_letter=letter,
        )
        ws.freeze_panes = "D5"

    def _driver_row(
        self,
        ws: Worksheet,
        cur: SheetCursor,
        label: str,
        revenue_row: int,
        driver_ref: str,
        letters: list[str],
        first_col: int,
        fmt: str,
        negate: bool = False,
    ) -> int:
        """Write `revenue x driver` for each forecast year."""
        row = cur.row
        ws.cell(row=row, column=1, value=label).font = S.LABEL_FONT
        sign = "-" if negate else ""
        for i, letter in enumerate(letters):
            driver = _cell(driver_ref, offset=i + 1)
            c = ws.cell(
                row=row, column=first_col + i, value=f"={sign}{letter}{revenue_row}*{driver}"
            )
            c.font, c.number_format = S.FORMULA_FONT, fmt
        cur.row += 1
        return row

    # -------------------------------------------------------------- sensitivity

    def _sensitivity(self, ws: Worksheet) -> None:
        """Trace each scenario through helper cells using the same terminal policy."""
        r = self.refs
        cfg = self.result.assumptions.sensitivity
        S.col_width(
            ws,
            {
                "A": 30,
                **{
                    get_column_letter(i): 16
                    for i in range(2, 2 + max(len(cfg.growth_deltas), len(cfg.multiple_deltas)))
                },
            },
        )
        cur = SheetCursor(ws, r)
        cur.title("Sensitivity Analysis")
        cur.note(
            "Only WACC and the named terminal input vary. Cost of equity for issuance stays fixed."
        )
        cur.note(
            "Trace every cell on SensitivityCalc. Unavailable cells are outside the model domain."
        )
        helper = ws.parent.create_sheet("SensitivityCalc")
        helper.append(
            [
                "Scenario",
                "WACC",
                "Perpetuity growth",
                "Exit multiple",
                "FCF5 Gordon",
                "Value driver",
                "Effective Gordon",
                "Exit value",
                "Method used",
                "Terminal value",
                "Equity value",
                "Value per share",
            ]
        )
        for col in range(1, 13):
            helper.column_dimensions[get_column_letter(col)].width = 21
            helper.cell(1, col).font = S.HEADER_FONT
            helper.cell(1, col).fill = S.HEADER_FILL
        first, last = r["first_letter"], r["last_letter"]
        fcf = f"DCF!${first}${r['fcf_row']}:${last}${r['fcf_row']}"
        exp = f"DCF!${first}${r['exp_row']}:${last}${r['exp_row']}"
        end = f"DCF!${last}${r['idx_row']}"
        mid = f"DCF!${last}${r['exp_row']}"
        for exit_grid in (False, True):
            cur.skip()
            cur.section("WACC against exit multiple" if exit_grid else "WACC against growth")
            header = cur.row
            ws.cell(header, 1, "WACC \\ multiple" if exit_grid else "WACC \\ growth")
            deltas = cfg.multiple_deltas if exit_grid else cfg.growth_deltas
            for j, delta in enumerate(deltas):
                formula = (
                    f"=MAX({r['exit_mult']}+({delta}),0.5)" if exit_grid else f"={r['g']}+({delta})"
                )
                ws.cell(header, 2 + j, formula).number_format = (
                    S.MULTIPLE if exit_grid else S.PERCENT_2
                )
            cur.row += 1
            for dw in cfg.wacc_deltas:
                row = cur.row
                ws.cell(row, 1, f"={r['wacc']}+({dw})").number_format = S.PERCENT_2
                for j in range(len(deltas)):
                    h = helper.max_row + 1
                    w = f"B{h}"
                    g = f"C{h}"
                    mv = f"D{h}"
                    gor = f"G{h}"
                    vd = f"F{h}"
                    ex = f"H{h}"
                    helper.cell(
                        h, 1, f"{'Exit' if exit_grid else 'Growth'} {get_column_letter(j + 2)}{row}"
                    )
                    helper.cell(h, 2, f"=Sensitivity!$A${row}")
                    helper.cell(
                        h,
                        3,
                        f"={r['g']}"
                        if exit_grid
                        else f"=Sensitivity!{get_column_letter(j + 2)}${header}",
                    )
                    helper.cell(
                        h,
                        4,
                        f"=Sensitivity!{get_column_letter(j + 2)}${header}"
                        if exit_grid
                        else f"={r['exit_mult']}",
                    )
                    helper.cell(h, 5, f"=IF({w}<={g},NA(),{r['term_fcf']}*(1+{g})/({w}-{g}))")
                    helper.cell(
                        h,
                        6,
                        f"=IF(OR({w}<={g},{r['term_nopat']}<=0),NA(),{r['term_nopat']}*(1+{g})*(1-{g}/IF({r['ronic']}>0,{r['ronic']},{w}))/({w}-{g}))",
                    )
                    # Exit-only fallback is ordinary Gordon, as in TerminalValue.compute.
                    mode = "FALSE" if exit_grid else f'{r["tv_pref"]}<>"exit_multiple"'
                    helper.cell(
                        h,
                        7,
                        f'=IF(AND({mode},{r["tv_mode"]}="value_driver",IFERROR({vd}>0,FALSE)),{vd},E{h})',
                    )
                    helper.cell(h, 8, f"=IF({r['term_ebitda']}<=0,NA(),{r['term_ebitda']}*{mv})")

                    def positive(ref):
                        return f"IFERROR({ref}>0,FALSE)"

                    normal = f'IF({positive(gor)},"gordon",IF({positive(ex)},"exit_multiple",IF({positive(vd)},"value_driver","gordon")))'
                    exfirst = f'IF({positive(ex)},"exit_multiple",{normal})'
                    vdfirst = f'IF({positive(vd)},"value_driver",{normal})'
                    method = (
                        exfirst
                        if exit_grid
                        else f'IF({r["tv_pref"]}="value_driver",{vdfirst},IF({r["tv_pref"]}="exit_multiple",{exfirst},{normal}))'
                    )
                    helper.cell(h, 9, "=" + method)
                    helper.cell(
                        h, 10, f'=IF(I{h}="gordon",{gor},IF(I{h}="exit_multiple",{ex},{vd}))'
                    )
                    helper.cell(
                        h,
                        11,
                        f'=SUMPRODUCT({fcf},1/((1+{w})^{exp}))+J{h}/(1+{w})^IF(I{h}="exit_multiple",{end},{mid})-{r["net_bridge"]}',
                    )
                    k = f'K{h}-IF({r["sbc_method"]}="dilute",{r["sbc_pv"]},0)'
                    domain = f"OR({w}<=0,{w}>1,{g}<-.02,{g}>.06,{mv}<=0,{mv}>100)"
                    if not exit_grid:
                        domain = f"OR({domain},{w}<={g})"
                    helper.cell(
                        h,
                        12,
                        f'=IF(OR({domain},AND({r["sbc_method"]}="dilute",{k}<=0)),NA(),({k})/({r["shares_in"]}+{r["overhang"]}))',
                    )
                    ws.cell(
                        row, j + 2, f"=SensitivityCalc!$L${h}"
                    ).number_format = self.price_format
                    for c in range(2, 13):
                        helper.cell(h, c).font = S.FORMULA_FONT
                    for c in [2, 3]:
                        helper.cell(h, c).number_format = S.PERCENT_2
                    for c in [4]:
                        helper.cell(h, c).number_format = S.MULTIPLE
                    for c in [5, 6, 7, 8, 10, 11]:
                        helper.cell(h, c).number_format = S.MONEY_MM
                    helper.cell(h, 12).number_format = self.price_format
                cur.row += 1
        helper.freeze_panes = "E2"
        ws.freeze_panes = "B4"

    # -------------------------------------------------------------------- comps

    def _comps(self, ws: Worksheet) -> None:
        S.col_width(ws, {"A": 14, "B": 18, "C": 18, "D": 14, "E": 14, "F": 14, "G": 14, "H": 16})
        cur = SheetCursor(ws, self.refs)
        cur.title("Comparable Companies", width=8)

        if self.comps is None or self.comps.table.empty:
            cur.note("No comparable-company data was available for this run.")
            return

        cur.note(f"Peer set source: {self.comps.peer_source}")
        cur.skip()

        columns = [
            ("market_cap", "Market cap", S.MONEY_MM),
            ("enterprise_value", "Enterprise value", S.MONEY_MM),
            ("ev_ebitda", "EV/EBITDA", S.MULTIPLE),
            ("ev_ebit", "EV/EBIT", S.MULTIPLE),
            ("ev_revenue", "EV/Revenue", S.MULTIPLE),
            ("pe", "P/E", S.MULTIPLE),
            ("revenue_growth", "Revenue growth", S.PERCENT),
            ("ebitda_margin", "EBITDA margin", S.PERCENT),
        ]
        cur.headers(["Ticker"] + [label for _, label, _ in columns])

        for ticker, row in self.comps.table.iterrows():
            is_target = ticker == self.comps.target
            cell = ws.cell(row=cur.row, column=1, value=str(ticker))
            cell.font = S.TOTAL_FONT if is_target else S.LABEL_FONT
            for j, (key, _, fmt) in enumerate(columns):
                value = row.get(key)
                cell = ws.cell(
                    row=cur.row,
                    column=2 + j,
                    value=None if value is None or pd.isna(value) else float(value),
                )
                cell.font = S.TOTAL_FONT if is_target else S.FORMULA_FONT
                cell.number_format = fmt
            cur.row += 1

        cur.skip()
        cur.section("Peer medians (target excluded)", width=9)
        for key, label, fmt in columns[2:]:
            median = self.comps.medians.get(f"{key}_median")
            if median is None:
                continue
            cur.label_value(label, float(median), fmt, kind="formula")
        cur.label_value(
            "Screened peer count", float(self.comps.peer_count), S.INTEGER, kind="formula"
        )

        notes = self.comps.notes()
        if notes:
            cur.skip()
            cur.section("Notes", width=9)
            for note in notes:
                cur.note(note, warn=True)

    # ----------------------------------------------------------- football field

    def _football(self, ws: Worksheet) -> None:
        S.col_width(ws, {"A": 34, "B": 14, "C": 14, "D": 14, "E": 14})
        cur = SheetCursor(ws, self.refs)
        cur.title("Valuation Summary - Football Field", width=5)
        cur.note(
            "Ranges use the stated sensitivity or simulation. Single-point methods use a thin display band (0.4% total width, minimum 0.01); it is not an uncertainty estimate."
        )
        cur.skip()

        if self.football is None or self.football.empty:
            cur.note("No valuation ranges were produced for this run.")
            return

        cur.headers(["Method", "Low", "Span", "High", "Base case"])
        first_data = cur.row

        for _, row in self.football.iterrows():
            low = _num(row.get("low"))
            high = _num(row.get("high"))
            mid = _num(row.get("midpoint"))
            ws.cell(row=cur.row, column=1, value=str(row["method"])).font = S.LABEL_FONT
            for col, value, fmt in (
                (2, low, self.price_format),
                (3, None if low is None or high is None else high - low, self.price_format),
                (4, high, self.price_format),
                (5, mid, self.price_format),
            ):
                cell = ws.cell(row=cur.row, column=col, value=value)
                cell.font = S.FORMULA_FONT
                cell.number_format = fmt
            cur.row += 1

        last_data = cur.row - 1

        # A floating bar: an invisible base up to `low`, then a visible span.
        chart = BarChart()
        chart.type = "bar"
        chart.grouping = "stacked"
        chart.overlap = 100
        chart.title = "Implied value per share"
        chart.y_axis.title = f"{getattr(self.financials, 'currency', 'quote currency')} per share"
        chart.height, chart.width = 9, 18

        data = Reference(ws, min_col=2, max_col=3, min_row=first_data - 1, max_row=last_data)
        categories = Reference(ws, min_col=1, min_row=first_data, max_row=last_data)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        if chart.series:
            chart.series[0].graphicalProperties.noFill = True
            chart.series[0].graphicalProperties.line.noFill = True
        ws.add_chart(chart, f"A{last_data + 3}")

    # -------------------------------------------------------------- historicals

    def _historicals(self, ws: Worksheet) -> None:
        cur = SheetCursor(ws, self.refs)
        cur.title("Reported Historical Financials", width=8)
        cur.note("As reported by the data source, mapped to canonical line items.")
        cur.skip()

        if self.financials is None or self.financials.statements.empty:
            cur.note("No historical statements were loaded.")
            return

        statements = self.financials.statements
        S.col_width(ws, {"A": 34})
        cur.headers(
            ["Line item"]
            + [
                str(p)[:10] if not hasattr(p, "date") else p.date().isoformat()
                for p in statements.columns
            ]
        )
        for col in range(2, statements.shape[1] + 2):
            ws.column_dimensions[get_column_letter(col)].width = 16

        for name in statements.index:
            ws.cell(row=cur.row, column=1, value=name).font = S.LABEL_FONT
            for j, period in enumerate(statements.columns):
                value = statements.loc[name, period]
                cell = ws.cell(
                    row=cur.row, column=2 + j, value=None if pd.isna(value) else float(value)
                )
                cell.font = S.INPUT_FONT
                cell.number_format = S.MONEY_MM
            cur.row += 1

        ws.freeze_panes = "B5"

    # ------------------------------------------------------------------ summary

    def _summary(self, ws: Worksheet) -> None:
        S.col_width(ws, {"A": 46, "B": 20, "C": 20})
        cur = SheetCursor(ws, self.refs)
        r = self.refs
        result = self.result

        cur.title(f"{result.ticker} - Valuation Summary", width=4)
        cur.note(
            "Amounts and shares are in millions; prices are per share. Headline figures link to DCF. The reported-figure blocks below "
            "(comps, historicals and the SBC memo) are written as values, not formulas, "
            "so they will not follow a change made on Inputs -- rebuild the workbook."
        )
        cur.skip()

        cur.section("Conclusion", width=4)
        cur.label_value("Implied value per share", f"={r['vps']}", self.price_format, kind="link")
        cur.label_value("Snapshot share price", f"={r['price']}", self.price_format, kind="link")
        cur.label_value("Upside / (downside)", f"={r['upside']}", S.PERCENT, kind="link")
        cur.skip()

        cur.section("Valuation build", width=4)
        cur.label_value("Enterprise value", f"={r['ev']}", S.MONEY_MM, kind="link")
        cur.label_value("Equity value", f"={r['equity']}", S.MONEY_MM, kind="link")
        cur.label_value("Diluted shares", f"={r['shares_out']}", S.SHARES_MM, kind="link")
        cur.label_value("WACC", f"={r['wacc']}", S.PERCENT_2, kind="link")
        cur.label_value("Terminal value as % of EV", f"={r['tv_pct']}", S.PERCENT, kind="link")
        cur.skip()

        cur.section("Free cash flow definitions", width=4)
        reported = result.projection.reported_fcf_latest
        cur.label_value(
            "Reported FCF (CFO less capex, latest year)",
            None if pd.isna(reported) else float(reported),
            S.MONEY_MM,
            kind="input",
        )
        cur.label_value(
            "Adjusted FCF, year 1 (SBC expensed)",
            float(result.projection.adjusted_fcf.iloc[0]),
            S.MONEY_MM,
            kind="input",
        )
        cur.label_value(
            "SBC-neutral FCF, year 1 (SBC added back)",
            float(result.projection.sbc_neutral_fcf.iloc[0]),
            S.MONEY_MM,
            kind="input",
        )
        cur.label_value("SBC treatment used", result.assumptions.sbc.method, kind="input")
        cur.note(
            "Reported FCF treats stock compensation as free. This model does not: it is "
            "charged either against cash flow or against the share count, never both."
        )
        cur.skip()

        if self.financials is not None:
            try:
                from src.dcf.moat import analyze_moat
                from src.dcf.reverse_dcf import solve_reverse_dcf

                moat = analyze_moat(self.financials, result)
                rev = solve_reverse_dcf(
                    self.financials,
                    result.assumptions,
                    self.comps.terminal_inputs() if self.comps else {},
                    ticker=result.ticker,
                )

                cur.section("Accounting Capital Efficiency (Moat Unverified)", width=4)
                cur.label_value(
                    "Invested capital (base)", moat.invested_capital_base, S.MONEY_MM, kind="input"
                )
                cur.label_value("Base ROIC", moat.roic_base, S.PERCENT_2, kind="input")
                cur.label_value(
                    "Economic spread (ROIC - WACC)", moat.economic_spread, S.PERCENT_2, kind="input"
                )
                cur.label_value("Moat assessment", moat.moat_rating, kind="input")
                cur.skip()

                cur.section("Market Expectations & Margin of Safety", width=4)
                if rev.implied_revenue_growth_cagr is not None:
                    cur.label_value(
                        "Market-implied 5-yr revenue CAGR",
                        rev.implied_revenue_growth_cagr,
                        S.PERCENT_2,
                        kind="input",
                    )
                else:
                    cur.label_value(
                        "Market-implied 5-yr revenue CAGR", rev.implied_revenue_status, kind="input"
                    )

                if rev.implied_perpetuity_growth is not None:
                    cur.label_value(
                        "Market-implied perpetuity growth (g)",
                        rev.implied_perpetuity_growth,
                        S.PERCENT_2,
                        kind="input",
                    )

                vps_ref = r["vps"]
                cur.label_value(
                    "Illustrative 15% discount to modeled value",
                    f"={vps_ref}*0.85",
                    self.price_format,
                    kind="link",
                )
                cur.label_value(
                    "Illustrative 25% discount to modeled value",
                    f"={vps_ref}*0.75",
                    self.price_format,
                    kind="link",
                )
                cur.label_value(
                    "Illustrative 35% discount to modeled value",
                    f"={vps_ref}*0.65",
                    self.price_format,
                    kind="link",
                )
                cur.skip()
            except (ValueError, ArithmeticError) as exc:
                cur.note(f"Diagnostic unavailable: {exc}", warn=True)

        if self.monte_carlo:
            cur.section("Monte Carlo (WACC, terminal growth, EBIT margin)", width=4)
            for key, label, fmt in (
                ("p10", "10th percentile", self.price_format),
                ("p50", "Median", self.price_format),
                ("p90", "90th percentile", self.price_format),
                ("prob_above_market", "Probability above market price", S.PERCENT),
            ):
                if key in self.monte_carlo:
                    cur.label_value(label, float(self.monte_carlo[key]), fmt, kind="input")
            cur.skip()

        if result.warnings:
            cur.section("Model warnings", width=4)
            for message in result.warnings:
                cur.note(message, warn=True)

        ws.freeze_panes = "A3"


# ------------------------------------------------------------------- utilities


def _cell(range_ref: str, offset: int) -> str:
    """Shift a stored `Sheet!$C$12` reference `offset` columns to the right."""
    sheet, _, addr = range_ref.partition("!")
    column_letter = addr.split("$")[1]
    row = addr.split("$")[2]
    from openpyxl.utils import column_index_from_string

    new_col = get_column_letter(column_index_from_string(column_letter) + offset)
    return f"{sheet}!${new_col}${row}"


def _num(value: Any) -> float | None:
    """Coerce to a float Excel can hold.

    NaN becomes None, because NaN means "not reported" everywhere else in this
    codebase and a blank cell is the honest rendering of that.

    An infinity raises. It is never a legitimate figure in a valuation -- it means
    something upstream divided by zero -- and the alternatives are both worse than an
    error. Writing it produces `<c t="n"><v></v></c>`, a numeric cell holding nothing,
    which Excel may offer to repair. Quietly blanking it, which this function used to
    do, leaves a Summary line empty and every formula pointing at it reading zero:
    a confidently-presented wrong number with no warning, which is the exact failure
    mode this model is built to refuse.
    """
    if value is None or pd.isna(value):
        return None
    if not isinstance(value, (int, float, np.integer, np.floating)):
        return value
    number = float(value)
    if math.isinf(number):
        raise ValueError(
            "an infinite value reached the workbook, which means something upstream "
            "divided by zero. Excel cannot hold it and blanking it would leave a cell "
            "that reads as zero to every formula pointing at it."
        )
    return number


def _latest_period(financials: Any) -> str:
    """Period end of the most recent statement, for the "data as of" line."""
    statements = getattr(financials, "statements", None)
    if statements is None or not len(getattr(statements, "columns", [])):
        return "n/a"
    latest = statements.columns[-1]
    return latest.date().isoformat() if hasattr(latest, "date") else str(latest)[:10]


def assert_all_cells_finite(book: Workbook) -> None:
    """Fail loudly if any cell in an unsaved workbook holds a non-finite number.

    This is the only point at which the check is possible. `_num` sanitises the values
    that pass through it, but it guarded three call sites out of roughly forty when
    the second audit ran -- every other numeric write went straight to `ws.cell`. The
    two tests that claimed to cover this asserted on a *reloaded* workbook, where
    openpyxl has already replaced every non-finite value with None, so the guard
    `isinstance(cell.value, float)` was False and the assertion never executed.
    """
    offenders = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, float) and not math.isfinite(value):
                    offenders.append(f"{sheet.title}!{cell.coordinate} = {value}")
    if offenders:
        raise ValueError(
            "non-finite values would be written to the workbook, where they become "
            f"blank numeric cells no reader can distinguish from a real gap: {offenders[:10]}"
        )


def build_excel_model(
    result: Any,
    path: str | Path,
    financials: Any = None,
    comps: Any = None,
    monte_carlo: dict[str, float] | None = None,
    football: pd.DataFrame | None = None,
) -> Path:
    """Build the workbook.

    The Sensitivity sheet is generated from `result` and its assumptions rather than
    from a precomputed table: its cells are live Excel formulas, not pasted values, so
    a Python-side table would have nothing to contribute. The old `sensitivity` and
    `exit_sensitivity` parameters were accepted and silently discarded, which cost the
    CLI ~50 engine re-runs per model for nothing.
    """
    builder = ExcelModelBuilder(
        result=result,
        financials=financials,
        comps=comps,
        monte_carlo=monte_carlo,
        football=football,
    )
    return builder.build(path)
